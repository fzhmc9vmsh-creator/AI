#!/usr/bin/env python3
"""資料の文言の検査(audit_deck.py の配置・文字の収まりの検査と対で使う)。誤りが 1 件でもあれば FAIL(終了コード 1)。

    python3 lint_deck.py src [.] [--only s08]                  # ソースの検査(要所の数字は台帳から引いているか)
    python3 lint_deck.py pptx out/deck.pptx [--facts facts_deck.json]   # 組み上がった資料の照合

要所の数字は台帳から引く(src で検査):
  meta.js のタイトル、L.kpiTiles(…) の札、L.takeaway(…) の結論の箱、関係図の reason(評価の理由)にある数字は
  F('id') で引く。手で打った数字は誤り。どうしても例外にする行は末尾に  // lint-ok: 理由  と書く。
  その他(表・注記・図の中)の数字は手で書いてよい。代わりに pptx の照合で 3 点を機械で確かめる。

組み上がった資料の照合(pptx):
  1. 価値評価の数字: 「理論株価・期待値・標準偏差・適正株価・即時・分割・エントリー」の直後の数字が、
     valuate.py の表示値か台帳の表示値のどれかと一致するか(23.8 を 23.9 と書く誤りを止める)
  2. 古い時点の値: 台帳で資料の基準(as_of)より前の四半期の値が、時点ラベル(Q2・前期・前年・→ など)なしで出ていないか
  3. 未確認の値: 台帳で資料に使えない(not_for_deck、T4)とした値が出ていないか
  ほかに、版の言葉(第N版・前回・今回・v2・改訂)、仮置きの文字、計算の失敗(NaN・undefined)を止める。
"""
import glob
import json
import os
import re
import sys

NUM_UNIT = re.compile(r"(?<![A-Za-z_\-./])[−\-+]?\d[\d,]*(\.\d+)?\s?(%|％|倍|ドル|円|億|万|兆|pt|bp|bps|人|件|社)(?![A-Za-z])")
DECIMAL = re.compile(r"(?<![A-Za-z_\-./\d§])\d+\.\d+(?![\d.]*[A-Za-z_-])")
NUM = re.compile(r"(?<![\d.])\d[\d,]*(?:\.\d+)?(?![\d.])")
BANNED = [
    (re.compile(r"第\s?\d+\s?版"), "版の言葉(第N版)"), (re.compile(r"前回|今回"), "版の言葉(前回・今回)"),
    (re.compile(r"(?<![A-Za-z0-9_])v\d(?![\d.])"), "版の言葉(v2 など)"), (re.compile(r"改訂|新規スライド|差し替え"), "版の言葉(改訂など)"),
    (re.compile(r"TODO|TBD|XXX|〇〇|●●|△△|ダミー|仮置き"), "仮置きの文字"),
    (re.compile(r"undefined|NaN|\[object Object\]|(?<![A-Za-z])null(?![A-Za-z])"), "計算・参照の失敗"),
]
VAL_KEYS = re.compile(r"理論株価|期待値|標準偏差|適正株価|即時|分割|エントリー")
PERIOD_NEAR = re.compile(r"Q[1-4]|[1-4]Q|前期|前年|昨年|前四半期|(19|20)\d\d|→|から|時点|当時")
STR_RE = re.compile(r"'((?:[^'\\\n]|\\.)*)'|\"((?:[^\"\\\n]|\\.)*)\"|`((?:[^`\\]|\\.)*)`", re.S)
KEY_CALLS = ("kpiTiles(", "takeaway(")


# ---------------------------------------------------------------- ソース
def split_code(line):
    """1 行をコードとコメントに分ける(文字列の中の // はコメントにしない)。"""
    q = None
    i = 0
    while i < len(line):
        ch = line[i]
        if q:
            if ch == "\\":
                i += 2
                continue
            if ch == q:
                q = None
        elif ch in "'\"`":
            q = ch
        elif line[i:i + 2] == "//":
            return line[:i], line[i:]
        i += 1
    return line, ""


def paren_delta(code):
    """文字列の外の ( と ) の差。"""
    d, q, i = 0, None, 0
    while i < len(code):
        ch = code[i]
        if q:
            if ch == "\\":
                i += 2
                continue
            if ch == q:
                q = None
        elif ch in "'\"`":
            q = ch
        elif ch == "(":
            d += 1
        elif ch == ")":
            d -= 1
        i += 1
    return d


FIXED_TERMS = re.compile(r"PEG ?1(\.0)?(?![\d.])|PEG ?< ?1|S&P ?500|ルール・オブ・40|Rule of 40|1 株当たり|第\s?\d+\s?章")


def hard_numbers(code):
    code = re.sub(r"\$\{[^}]*\}", "", code)
    code = FIXED_TERMS.sub("", code)                 # 決まった言い回しの中の数字(PEG 1.0 など)は数えない
    code = re.sub(r"F(\.[nv])?\(\s*'[^']*'[^)]*\)", "", code)
    hits = []
    for m in STR_RE.finditer(code):
        s = next(g for g in m.groups() if g is not None)
        if re.fullmatch(r"[0-9A-Fa-f]{6}", s):
            continue
        for h in list(NUM_UNIT.finditer(s)) + list(DECIMAL.finditer(s)):
            t = h.group(0).strip()
            if not re.fullmatch(r"(19|20)\d\d", t):
                hits.append(t)
    return sorted(set(hits))


def lint_src(root, only=None):
    files = sorted(glob.glob(os.path.join(root, "slides", "*.js")))
    if only:
        files = [f for f in files if os.path.basename(f) == only + ".js"]
    elif os.path.exists(os.path.join(root, "meta.js")):
        files.append(os.path.join(root, "meta.js"))
    errs, info = [], []
    n_free = 0
    for fp in files:
        name = os.path.relpath(fp, root)
        is_meta = os.path.basename(fp) == "meta.js"
        depth, in_reason = 0, False
        with open(fp, encoding="utf-8") as fh:
            for n, raw in enumerate(fh, 1):
                code, comment = split_code(raw.rstrip("\n"))
                if not code.strip():
                    continue
                ok = "lint-ok" in comment
                if ok and not re.search(r"lint-ok\s*:\s*\S", comment):
                    errs.append(f"{name}:{n}: lint-ok に理由がない")
                if "noPeriod" in code:
                    info.append(f"{name}:{n}: 時点ラベルを外した数字がある(本文に時点が書いてあるか画像で確かめる)")
                # 要所の範囲: meta のタイトル / kpiTiles・takeaway の呼び出しの中 / reason の値
                starts_call = any(k in code for k in KEY_CALLS)
                if re.search(r"\breason\s*:", code):
                    in_reason = True
                elif in_reason and (re.match(r"\s*\w+\s*:", code) or re.match(r"\s*[}\]]", code)):
                    in_reason = False
                key = (is_meta and re.search(r"\btitle\s*:", code)) or depth > 0 or starts_call or in_reason
                if starts_call or depth > 0:
                    depth = max(0, depth + paren_delta(code))
                hits = hard_numbers(code)
                for rx, what in BANNED[:5]:
                    for m in STR_RE.finditer(code):
                        s = next(g for g in m.groups() if g is not None)
                        if rx.search(s) and not re.search(r"\bsource\s*:", code):
                            errs.append(f"{name}:{n}: {what}「{rx.search(s).group(0)}」")
                if not hits or ok:
                    continue
                if key:
                    where = "タイトル" if is_meta else "KPI の札・結論の箱・評価の理由"
                    errs.append(f"{name}:{n}: {where}の数字を手で打っている {hits[:4]} → F('id') で台帳から引く")
                else:
                    n_free += len(hits)
    return files, errs, info, n_free


# ---------------------------------------------------------------- 組み上がった資料
def slide_texts(path):
    from pptx import Presentation
    prs = Presentation(path)
    out = []
    for sl in prs.slides:
        t = []
        for sh in sl.shapes:
            if sh.has_text_frame:
                t.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                t += [c.text for r in sh.table.rows for c in r.cells]
        out.append("\n".join(t))
    return out


def qkey(period):
    m = re.fullmatch(r"Q([1-4]) (\d{4})", period or "")
    return (int(m.group(2)), int(m.group(1))) if m else None


def first_num(s):
    m = NUM.search(s or "")
    return m.group(0) if m else None


def value_rx(f):
    """台帳の値を資料の中で探す正規表現。数字と単位の組で探す(数字だけだと別の意味の数字に当たる)。
    有効数字が 1 桁の値(例: 3 億ドル)は誤検知が多いので、match_terms(近くにあるはずの語)が無ければ探さない。"""
    tok = first_num(f.get("disp"))
    if not tok:
        return None, None
    unit = (f.get("unit") or "").strip()
    sig = len(tok.replace(".", "").replace(",", "").lstrip("0"))
    terms = f.get("match_terms") or []
    if sig < 2 and not terms:
        return None, None
    u = re.escape(unit[:1]) if unit and unit not in ("%",) else (r"%" if unit == "%" else "")
    rx = re.compile(r"(?<![\d.])" + re.escape(tok) + (r"\s?" + u if u else "") + r"(?![\d.])")
    return rx, terms


def lint_pptx(path, facts_path):
    texts = slide_texts(path)
    errs, notes = [], []
    for i, t in enumerate(texts, 1):
        for rx, what in BANNED:
            for m in rx.finditer(t):
                errs.append(f"p.{i}: {what}「{m.group(0)}」… {t[max(0, m.start() - 15):m.end() + 15]!r}")
    if not facts_path or not os.path.exists(facts_path):
        notes.append("facts_deck.json が無いため、数字の照合(価値評価・古い時点・未確認)は行っていない")
        return len(texts), errs, notes
    with open(facts_path, encoding="utf-8") as fh:
        db = json.load(fh)
    facts = db.get("facts", {})
    allowed = set()
    for f in facts.values():
        allowed.update(NUM.findall(str(f.get("disp", ""))))
    # 1. 価値評価の数字
    for i, t in enumerate(texts, 1):
        for m in VAL_KEYS.finditer(t):
            # 語の直後(区切り「。、・改行」をまたがず 10 文字以内)の、価格を表す数字だけを見る(+0.5pt・16 通りなどは見ない)
            tail = t[m.end():m.end() + 40]
            nm = re.match(r"[^\d\n。、・]{0,10}?(?<![A-Za-z\d.+−-])(\d[\d,]*(?:\.\d+)?)(?![\d.])(?!\s?(pt|%|bp|通り|倍|年|件|つ|か月|四半期))", tail)
            if nm and nm.group(1) not in allowed:
                nm = type("M", (), {"group": lambda self, k=0, v=nm.group(1): v})()
            else:
                nm = None
            if nm and not re.match(r"\s*[(（]?\s*(19|20)\d\d", tail):
                errs.append(f"p.{i}: 「{m.group(0)}」の数字 {nm.group(0)} が計算機・台帳のどの値とも一致しない … {t[m.start():m.end() + 16]!r}")
    # 2. 古い時点の値
    now = qkey(db.get("as_of", ""))
    if now:
        current = set()
        for f in facts.values():
            if qkey(f.get("period")) == now:
                current.update(NUM.findall(str(f.get("disp", ""))))
        for fid, f in facts.items():
            k = qkey(f.get("period"))
            if not k or k >= now or f.get("stale_check") is False:
                continue
            tok = first_num(f.get("disp"))
            rx, terms = value_rx(f)
            if not rx or tok in current:
                continue
            for i, t in enumerate(texts, 1):
                for m in rx.finditer(t):
                    # 時点のラベルは同じ行(表なら同じセル)の中で探す
                    ls, le = t.rfind("\n", 0, m.start()) + 1, t.find("\n", m.end())
                    le = len(t) if le < 0 else le
                    near = t[max(ls, m.start() - 20):min(le, m.end() + 14)]
                    if terms and not any(w in t[max(0, m.start() - 30):m.end() + 30] for w in terms):
                        continue
                    if not PERIOD_NEAR.search(near):
                        errs.append(f"p.{i}: 古い時点({f.get('period')})の値 {tok}({fid})が時点ラベルなしで出ている … {near!r}")
    # 3. 未確認の値
    for f in db.get("forbidden", []):
        tok = first_num(f.get("disp"))
        rx, terms = value_rx(f)
        if not rx:
            notes.append(f"{f.get('id')}: 有効数字 1 桁の値は match_terms が無いと探せない(台帳に近くの語を足す)")
            continue
        for i, t in enumerate(texts, 1):
            for m in rx.finditer(t):
                near = t[max(0, m.start() - 15):m.end() + 10]
                if terms and not any(w in t[max(0, m.start() - 30):m.end() + 30] for w in terms):
                    continue
                errs.append(f"p.{i}: 資料に使えない値({f.get('id')}、{f.get('reason', '未確認')})と同じ数字 {tok} … {near!r}")
    return len(texts), errs, notes


def main():
    a = sys.argv[1:]
    if not a or a[0] not in ("src", "pptx"):
        print(__doc__)
        sys.exit(2)
    only = facts = None
    if "--only" in a:
        only = a[a.index("--only") + 1]
        a = [x for x in a if x not in ("--only", only)]
    if "--facts" in a:
        facts = a[a.index("--facts") + 1]
        a = [x for x in a if x not in ("--facts", facts)]
    if a[0] == "src":
        files, errs, info, n_free = lint_src(a[1] if len(a) > 1 else ".", only)
        for x in info:
            print("確認:", x)
        for e in errs:
            print("誤り:", e)
        print(f"{len(files)} ファイル / 誤り {len(errs)} / 確認 {len(info)} / 要所の外の手書きの数字 {n_free} 個(可。pptx の照合で確かめる)"
              f" → {'PASS' if not errs else 'FAIL'}")
    else:
        path = a[1]
        if facts is None:
            for cand in (os.path.join(os.path.dirname(os.path.abspath(path)), "..", "facts_deck.json"), "facts_deck.json"):
                if os.path.exists(cand):
                    facts = cand
                    break
        n, errs, notes = lint_pptx(path, facts)
        for x in notes:
            print("注意:", x)
        for e in errs:
            print("誤り:", e)
        print(f"{n} 枚 / 誤り {len(errs)} → {'PASS' if not errs else 'FAIL'}")
    sys.exit(1 if errs else 0)


if __name__ == "__main__":
    main()
