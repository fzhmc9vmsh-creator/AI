# -*- coding: utf-8 -*-
"""DESIGN_SPEC.md §3 の部品ライブラリ (python-pptx 1.0.2)。

数値・色は仕様書 §0〜§2 の写経。座標・寸法はすべて inch。
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_MARKER_STYLE
from pptx.chart.data import CategoryChartData
from pptx.oxml.ns import qn

# ---- §1 カラーパレット ----------------------------------------------------
PRIMARY = "197A56"
PRIMARY_DARK = "0E4D33"
SECONDARY = "86C99B"
ACCENT = "1BA1A6"
NEG = "C0392B"
AMBER = "D68910"
TEXT = "333333"
SUB = "767676"
RULE = "D9D9D9"
WHITE = "FFFFFF"
LIGHT_BG = "EEF6F1"
GRAY_BG = "F4F4F4"
BADGE_ORANGE = "E67E22"    # バッジ専用
BADGE_DARKRED = "7B241C"   # バッジ専用

# 評価バッジ: grade -> (fill, text または None=自動でコントラストの高い方)
BADGE_COLORS = {
    "S": (PRIMARY_DARK, None),
    "A": (PRIMARY, None),
    "B": (SECONDARY, PRIMARY_DARK),
    "C": (AMBER, None),
    "D": (BADGE_ORANGE, None),
    "E": (NEG, None),
    "F": (BADGE_DARKRED, None),
    "NA": (SUB, None),
}

SERIES_COLORS = [PRIMARY, SECONDARY, ACCENT, SUB]  # §3 グラフ系列色の順
FONT = "Meiryo UI"
SLIDE_W = 13.333
SLIDE_H = 7.5
INNER_MARGIN = 0.05
NO_STYLE_NO_GRID = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # 表スタイル無効

_ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
_VANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}
_FONT_SUCCESSORS = {qn(t) for t in ("a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")}


# ---- 色・フォント内部ヘルパ ---------------------------------------------------
def _rgb(hexstr):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(hexstr)


def _lin(c255):
    c = c255 / 255.0
    return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4


def _lum(hexstr):
    r, g, b = (int(hexstr[i:i + 2], 16) for i in (0, 2, 4))
    return 0.2126 * _lin(r) + 0.7152 * _lin(g) + 0.0722 * _lin(b)


def _auto_text(fill):
    """塗り色の上で読みやすい文字色(WHITE か TEXT)を選ぶ。"""
    L = _lum(fill)
    contrast_white = (1.05) / (L + 0.05)
    contrast_text = (L + 0.05) / (_lum(TEXT) + 0.05)
    return WHITE if contrast_white >= contrast_text else TEXT


def _enforce_font(rpr):
    """a:rPr / a:defRPr に latin と ea を Meiryo UI で入れる(スキーマ順を守る)。"""
    for tag in ("a:latin", "a:ea"):
        for el in rpr.findall(qn(tag)):
            rpr.remove(el)
    latin = rpr.makeelement(qn("a:latin"), {"typeface": FONT})
    ea = rpr.makeelement(qn("a:ea"), {"typeface": FONT})
    succ = next((c for c in rpr if c.tag in _FONT_SUCCESSORS), None)
    if succ is None:
        rpr.append(latin)
        rpr.append(ea)
    else:
        succ.addprevious(latin)
        succ.addprevious(ea)


def _style_run(run, size, bold=False, color=TEXT):
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = _rgb(color)
    _enforce_font(run._r.get_or_add_rPr())


def _write_paragraphs(tf, text, size, bold, color, align):
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = _ALIGN[align]
        run = p.add_run()
        run.text = line
        _style_run(run, size, bold, color)


def _fill_tf(tf, text, size, bold=False, color=TEXT, align="l", valign="t"):
    """テキストボックス系の共通設定: 折返し・auto_size なし・内側余白 0.05in。"""
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(INNER_MARGIN)
    tf.margin_top = tf.margin_bottom = Inches(INNER_MARGIN)
    tf.vertical_anchor = _VANCHOR[valign]
    _write_paragraphs(tf, text, size, bold, color, align)


def _plain(shape):
    """python-pptx 既定の p:style(テーマ色参照)を外す。色は全部明示する。"""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)
    return shape


def _set_fill(shape, fill):
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill)


def _set_line(shape, line, width_pt=0.75):
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = _rgb(line)
        shape.line.width = Pt(width_pt)


def _rect(slide, x, y, w, h, name, shape_type=MSO_SHAPE.RECTANGLE,
          fill=None, line=None, line_w=0.75):
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    _plain(s)
    _set_fill(s, fill)
    _set_line(s, line, line_w)
    s.name = name
    return s


def _text_at(slide, x, y, w, h, text, size, bold=False, color=TEXT,
             align="l", valign="t", name="text"):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    _fill_tf(s.text_frame, text, size, bold, color, align, valign)
    s.name = name
    return s


# ---- スライド・プレゼン ----------------------------------------------------
def new_prs():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    return prs


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # 白紙


def set_type(slide, t):
    """notes に type=cover|divider|closing|content を書く(audit.py が読む)。"""
    if t not in ("cover", "divider", "closing", "content"):
        raise ValueError(f"unknown type: {t}")
    slide.notes_slide.notes_text_frame.text = f"type={t}"


# ---- §3 部品 ---------------------------------------------------------------
def add_title(slide, title, kicker=None):
    t = _text_at(slide, 0.5, 0.25, 12.3, 0.78, title, 20, bold=True,
                 color=PRIMARY, valign="b", name="title")
    if kicker:
        _text_at(slide, 0.5, 1.10, 12.3, 0.3, kicker, 11, color=SUB, name="kicker")
    return t


def add_source(slide, text):
    if not text.startswith("出所:"):
        text = "出所: " + text
    return _text_at(slide, 0.5, 6.95, 10.8, 0.35, text, 8, color=SUB, name="source")


def add_page_number(slide, n):
    return _text_at(slide, 12.2, 6.98, 0.65, 0.3, str(n), 8, color=SUB,
                    align="r", name="page_number")


def badge(slide, x, y, grade, size=0.42):
    if grade not in BADGE_COLORS:
        raise ValueError(f"unknown grade: {grade}")
    fill, fg = BADGE_COLORS[grade]
    fg = fg or _auto_text(fill)
    s = _rect(slide, x, y, size, size, f"badge_{grade}",
              shape_type=MSO_SHAPE.ROUNDED_RECTANGLE, fill=fill, line=None)
    _fill_tf(s.text_frame, grade, round(size * 33), bold=True, color=fg,
             align="c", valign="m")
    return s


def rank_badge(slide, x, y, rank, d=0.42):
    s = _rect(slide, x, y, d, d, f"rank_{rank}", shape_type=MSO_SHAPE.OVAL,
              fill=PRIMARY_DARK, line=None)
    _fill_tf(s.text_frame, str(rank), 14, bold=True, color=WHITE,
             align="c", valign="m")
    return s


def box(slide, x, y, w, h, fill=None, line=RULE, text=None, size=11, bold=False,
        color=TEXT, align="l", valign="t", name="box"):
    s = _rect(slide, x, y, w, h, name, fill=fill, line=line)
    if text is not None:
        _fill_tf(s.text_frame, text, size, bold, color, align, valign)
    return s


def takeaway(slide, x, y, w, h, text):
    _rect(slide, x, y, w, h, "bg_takeaway", fill=LIGHT_BG, line=None)
    _rect(slide, x, y, 0.06, h, "bg_takeaway_bar", fill=PRIMARY, line=None)
    return _text_at(slide, x + 0.2, y, w - 0.3, h, text, 11, color=TEXT,
                    valign="m", name="takeaway")


def table(slide, x, y, w, h, rows, col_widths, header=True, font=10.5,
          header_fill=PRIMARY_DARK, header_color=WHITE, zebra=False,
          row_height=None, cell_colors=None, bold_cells=None, col_align=None,
          name="table"):
    """表。縦罫なし・水平線 RULE のみ。cell_colors は {(r,c): hex} のセル塗り。"""
    cell_colors = cell_colors or {}
    bold_cells = bold_cells or set()
    nr, nc = len(rows), len(rows[0])
    if any(len(r) != nc for r in rows):
        raise ValueError("rows の列数が揃っていない")
    if len(col_widths) != nc:
        raise ValueError("col_widths の数が列数と違う")
    if abs(sum(col_widths) - w) > 0.01:
        raise ValueError(f"col_widths の合計 {sum(col_widths)} != w {w}")
    col_align = col_align or ["l"] * nc

    if row_height is None:
        heights = [h / nr] * nr
    elif isinstance(row_height, (int, float)):
        heights = [row_height] * nr
    else:
        heights = list(row_height)

    gf = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(sum(heights)))
    gf.name = name
    tbl = gf.table

    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = tblPr.makeelement(qn("a:tableStyleId"), {})
        tblPr.append(sid)
    sid.text = NO_STYLE_NO_GRID

    for c, cw in enumerate(col_widths):
        tbl.columns[c].width = Inches(cw)
    for r, rh in enumerate(heights):
        tbl.rows[r].height = Inches(rh)

    for r in range(nr):
        is_header = header and r == 0
        for c in range(nc):
            cell = tbl.cell(r, c)
            _cell_borders(cell)  # 塗りより前に ln を入れる(スキーマ順)
            if (r, c) in cell_colors:
                fill = cell_colors[(r, c)]
                fg = _auto_text(fill)
            elif is_header:
                fill, fg = header_fill, header_color
            elif zebra and r % 2 == 0:
                fill, fg = GRAY_BG, TEXT
            else:
                fill, fg = WHITE, TEXT
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(fill)
            cell.margin_left = cell.margin_right = Inches(INNER_MARGIN)
            cell.margin_top = cell.margin_bottom = Inches(INNER_MARGIN)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            bold = is_header or (r, c) in bold_cells
            _write_paragraphs(tf, rows[r][c], font, bold, fg, col_align[c])

    gf.height = Inches(sum(heights))
    return gf


def _cell_borders(cell):
    """縦罫なし(lnL/lnR=noFill)、水平線のみ RULE(lnT/lnB)。"""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
    specs = [("a:lnL", None), ("a:lnR", None), ("a:lnT", RULE), ("a:lnB", RULE)]
    for idx, (tag, color) in enumerate(specs):
        ln = tcPr.makeelement(qn(tag), {"w": str(int(Pt(0.75)))})
        if color is None:
            ln.set("w", "0")
            ln.append(ln.makeelement(qn("a:noFill"), {}))
        else:
            sf = ln.makeelement(qn("a:solidFill"), {})
            sf.append(sf.makeelement(qn("a:srgbClr"), {"val": color}))
            ln.append(sf)
        tcPr.insert(idx, ln)


def arrow(slide, x1, y1, x2, y2, color=SUB, width_pt=1.5, dashed=False, head=True,
          name="arrow"):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    _plain(c)
    c.line.color.rgb = _rgb(color)
    c.line.width = Pt(width_pt)
    if dashed:
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if head:
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    c.name = name
    return c


def chevron(slide, x, y, w, h, text, fill):
    s = _rect(slide, x, y, w, h, "chevron", shape_type=MSO_SHAPE.CHEVRON,
              fill=fill, line=None)
    _fill_tf(s.text_frame, text, 11, bold=True, color=_auto_text(fill),
             align="c", valign="m")
    return s


def kpi_card(slide, x, y, w, h, value, label, note=None, color=PRIMARY):
    if h < 1.5:
        raise ValueError("kpi_card の高さは 1.5in 以上")
    _rect(slide, x, y, w, h, "bg_kpi_card", fill=WHITE, line=RULE)
    _text_at(slide, x + 0.2, y + 0.08, w - 0.4, 0.85, value, 40, bold=True,
             color=color, name="kpi_value")
    _text_at(slide, x + 0.2, y + 0.93, w - 0.4, 0.3, label, 11, color=TEXT,
             name="kpi_label")
    if note:
        _text_at(slide, x + 0.2, y + 1.2, w - 0.4, 0.26, note, 9, color=SUB,
                 name="kpi_note")


def section_divider(prs, chapter_no, chapters, current_idx):
    """濃緑背景のトラッカー型目次。chapter_no は現在章の 1 始まり番号で、current_idx(0 始まり)と一致必須。"""
    if not 0 <= current_idx < len(chapters):
        raise ValueError("current_idx が範囲外")
    if chapter_no != current_idx + 1:
        raise ValueError("chapter_no は current_idx+1 と一致させる")
    s = new_slide(prs)
    set_type(s, "divider")
    _rect(s, 0, 0, SLIDE_W, SLIDE_H, "bg_divider", fill=PRIMARY_DARK)
    _text_at(s, 0.9, 2.3, 3.2, 0.7, "目次", 30, bold=True, color=WHITE, name="toc_label")
    for i, ch in enumerate(chapters):
        y = 2.3 + i * 1.1
        cur = i == current_idx
        no = _rect(s, 4.6, y, 0.52, 0.52, f"toc_no_{i + 1}", shape_type=MSO_SHAPE.OVAL,
                   fill=WHITE if cur else None, line=None if cur else SECONDARY, line_w=1.25)
        _fill_tf(no.text_frame, str(i + 1), 14, bold=True,
                 color=PRIMARY_DARK if cur else SECONDARY, align="c", valign="m")
        _text_at(s, 5.4, y, 7.2, 0.52, ch, 20, bold=cur,
                 color=WHITE if cur else SECONDARY, valign="m", name=f"toc_item_{i + 1}")
    return s


def cover(prs, title, subtitle, date):
    s = new_slide(prs)
    set_type(s, "cover")
    _rect(s, 0, 0, SLIDE_W, SLIDE_H, "bg_base", fill=WHITE)
    _rect(s, 0, 0, 8.6, SLIDE_H, "bg_cover_panel", fill=PRIMARY_DARK)
    _text_at(s, 0.8, 2.2, 7.2, 1.6, title, 30, bold=True, color=WHITE,
             valign="b", name="cover_title")
    _text_at(s, 0.8, 4.0, 7.2, 0.6, subtitle, 14, color=WHITE, name="cover_subtitle")
    _text_at(s, 0.8, 4.7, 7.2, 0.4, date, 12, color=WHITE, name="cover_date")
    return s


def closing(prs, text):
    s = new_slide(prs)
    set_type(s, "closing")
    _rect(s, 0, 0, SLIDE_W, SLIDE_H, "bg_closing", fill=PRIMARY_DARK)
    _text_at(s, 1.0, 2.4, 11.333, 2.7, text, 28, bold=True, color=WHITE,
             align="c", valign="m", name="closing_text")
    return s


# ---- ネイティブチャート -------------------------------------------------------
def _style_chart(chart, number_format, legend):
    chart.has_title = False
    chart.font.name = FONT
    chart.font.size = Pt(9)
    chart.font.color.rgb = _rgb(TEXT)

    va = chart.value_axis
    va.has_major_gridlines = True
    gl = va.major_gridlines.format.line
    gl.color.rgb = _rgb(RULE)
    gl.width = Pt(0.5)
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(9)
    va.tick_labels.font.color.rgb = _rgb(SUB)
    va.tick_labels.number_format = number_format
    va.tick_labels.number_format_is_linked = False

    ca = chart.category_axis
    ca.has_major_gridlines = False
    ca.format.line.color.rgb = _rgb(RULE)
    ca.tick_labels.font.size = Pt(9)
    ca.tick_labels.font.color.rgb = _rgb(TEXT)

    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(9)
        chart.legend.font.color.rgb = _rgb(TEXT)


def _series_labels(series, number_format, position, color):
    dl = series.data_labels
    dl.show_value = True
    dl.number_format = number_format
    dl.number_format_is_linked = False
    dl.position = position
    dl.font.size = Pt(9)
    dl.font.color.rgb = _rgb(color)


def _finish_chart_fonts(chart):
    for el in chart._chartSpace.iter(qn("a:defRPr"), qn("a:rPr"), qn("a:endParaRPr")):
        _enforce_font(el)


def bar_chart(slide, x, y, w, h, categories, series_dict, stacked=False, horizontal=False,
              number_format="0", data_labels=True, colors=None, name="chart_bar"):
    names = list(series_dict.keys())
    colors = colors or SERIES_COLORS[:len(names)]
    if len(colors) < len(names):
        raise ValueError("colors が系列数より少ない")
    if horizontal:
        ct = XL_CHART_TYPE.BAR_STACKED if stacked else XL_CHART_TYPE.BAR_CLUSTERED
    else:
        ct = XL_CHART_TYPE.COLUMN_STACKED if stacked else XL_CHART_TYPE.COLUMN_CLUSTERED

    cd = CategoryChartData(number_format=number_format)
    cd.categories = list(categories)
    for n in names:
        cd.add_series(n, list(series_dict[n]))

    gf = slide.shapes.add_chart(ct, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    gf.name = name
    chart = gf.chart
    plot = chart.plots[0]
    plot.gap_width = 60
    if stacked:
        plot.overlap = 100
    _style_chart(chart, number_format, legend=len(names) > 1)
    if horizontal:
        chart.category_axis.reverse_order = True

    for s, c in zip(plot.series, colors):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = _rgb(c)
        s.format.line.fill.background()
        if data_labels:
            if stacked:
                _series_labels(s, number_format, XL_LABEL_POSITION.CENTER,
                               TEXT if c == SECONDARY else WHITE)
            else:
                _series_labels(s, number_format, XL_LABEL_POSITION.OUTSIDE_END, TEXT)
    _finish_chart_fonts(chart)
    return gf


def line_chart(slide, x, y, w, h, categories, series_dict, number_format="0",
               data_labels=True, colors=None, name="chart_line"):
    names = list(series_dict.keys())
    colors = colors or SERIES_COLORS[:len(names)]
    if len(colors) < len(names):
        raise ValueError("colors が系列数より少ない")

    cd = CategoryChartData(number_format=number_format)
    cd.categories = list(categories)
    for n in names:
        cd.add_series(n, list(series_dict[n]))

    gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y),
                                Inches(w), Inches(h), cd)
    gf.name = name
    chart = gf.chart
    plot = chart.plots[0]
    _style_chart(chart, number_format, legend=len(names) > 1)

    for s, c in zip(plot.series, colors):
        s.smooth = False
        s.format.line.color.rgb = _rgb(c)
        s.format.line.width = Pt(2.25)
        s.marker.style = XL_MARKER_STYLE.CIRCLE
        s.marker.size = 7
        s.marker.format.fill.solid()
        s.marker.format.fill.fore_color.rgb = _rgb(c)
        s.marker.format.line.color.rgb = _rgb(c)
        if data_labels:
            _series_labels(s, number_format, XL_LABEL_POSITION.ABOVE, TEXT)
    _finish_chart_fonts(chart)
    return gf
