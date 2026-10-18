# -*- coding: utf-8 -*-
"""DESIGN_SPEC.md §4 の機械監査。使い方: python3 audit.py deck.pptx

出力: 違反ごとに「スライド番号 / shape名 / 項目 / 詳細」。0件なら最後の行に AUDIT PASS。
"""

import re
import sys
import unicodedata

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn

SLIDE_W, SLIDE_H = 13.333, 7.5
EPS = 0.001
TEXT_TOL = 0.02  # overlap 判定の閾値(in)

# §1 パレット(独立に写経)
PALETTE = {
    "197A56", "0E4D33", "86C99B", "1BA1A6", "C0392B", "D68910", "333333",
    "767676", "D9D9D9", "FFFFFF", "EEF6F1", "F4F4F4", "E67E22", "7B241C",
}
BADGE_ONLY = {"E67E22", "7B241C"}

FONT = "Meiryo UI"
MIN_PT = 7.0
SPECIAL = {"cover", "divider", "closing"}
BANNED_TITLE_END = ("分析", "について", "一覧", "概要")


def IN(emu):
    return emu / 914400.0


def slide_type(slide):
    if not slide.has_notes_slide:
        return "content"
    tf = slide.notes_slide.notes_text_frame
    txt = tf.text if tf is not None else ""
    m = re.search(r"type\s*=\s*(cover|divider|closing)", txt)
    return m.group(1) if m else "content"


# ---- 文字量の推定 ---------------------------------------------------------
def _cw(ch, size_pt):
    unit = size_pt / 72.0  # 1em in inch
    return (1.0 if unicodedata.east_asian_width(ch) in ("W", "F") else 0.55) * unit


def measure(paras, avail_w):
    """paras: [(text, size_pt)]。文字単位の折返しで (高さin, 行数) を返す。行高は 1.25em。"""
    total_h, total_lines = 0.0, 0
    for text, size in paras:
        unit = size / 72.0
        lines, cur = 1, 0.0
        for ch in text:
            w = _cw(ch, size)
            if cur + w > avail_w + 1e-9 and cur > 0:
                lines += 1
                cur = w
            else:
                cur += w
        total_lines += lines
        total_h += lines * 1.25 * unit
    return total_h, total_lines


def tf_paras(tf):
    """段落ごとの (text, size_pt)。サイズ未指定を返り値の 2 つ目で報告する。"""
    out, missing = [], False
    for p in tf.paragraphs:
        text = "".join(r.text for r in p.runs)
        sizes = [r.font.size.pt for r in p.runs if r.font.size is not None]
        if sizes:
            size = max(sizes)
        else:
            size = 18.0
            if p.runs:
                missing = True
        out.append((text, size))
    return out, missing


def tf_margins(tf):
    return (IN(tf.margin_left), IN(tf.margin_right), IN(tf.margin_top), IN(tf.margin_bottom))


# ---- 個別チェック -----------------------------------------------------------
def check_bounds(sno, sh, v):
    x, y, w, h = IN(sh.left), IN(sh.top), IN(sh.width), IN(sh.height)
    if x < -EPS or y < -EPS or x + w > SLIDE_W + EPS or y + h > SLIDE_H + EPS:
        v.append((sno, sh.name, "bounds",
                  f"スライド外: x={x:.3f} y={y:.3f} 右端={x + w:.3f} 下端={y + h:.3f}"))


def check_palette(sno, sh, v):
    root = sh.chart._chartSpace if sh.has_chart else sh._element
    colors = {el.get("val").upper() for el in root.iter(qn("a:srgbClr"))}
    for c in sorted(colors):
        if c not in PALETTE:
            v.append((sno, sh.name, "palette", f"パレット外の色 #{c}"))
        elif c in BADGE_ONLY and not sh.name.startswith("badge_"):
            v.append((sno, sh.name, "palette", f"バッジ専用色 #{c} をバッジ以外に使用"))
    if any(True for _ in root.iter(qn("a:schemeClr"))):
        v.append((sno, sh.name, "palette", "テーマ色参照(schemeClr)が残っている"))


def _font_problems(tf):
    probs = set()
    for p in tf.paragraphs:
        for r in p.runs:
            if r.font.name != FONT:
                probs.add(f"フォント名={r.font.name}")
            rPr = r._r.rPr
            ea = rPr.find(qn("a:ea")) if rPr is not None else None
            if ea is None or ea.get("typeface") != FONT:
                probs.add("ea が Meiryo UI でない")
            if r.font.size is None:
                probs.add("サイズ未指定")
            elif r.font.size.pt < MIN_PT:
                probs.add(f"サイズ{r.font.size.pt}pt(<7pt)")
    return probs


def check_font(sno, sh, v):
    targets = []
    if sh.has_table:
        for row in sh.table.rows:
            for cell in row.cells:
                targets.append(cell.text_frame)
    elif sh.has_text_frame:
        targets.append(sh.text_frame)
    probs = set()
    for tf in targets:
        probs |= _font_problems(tf)
    if sh.has_chart:
        for el in sh.chart._chartSpace.iter(qn("a:defRPr"), qn("a:rPr")):
            lat = el.find(qn("a:latin"))
            if lat is not None and lat.get("typeface") != FONT:
                probs.add(f"グラフのフォント={lat.get('typeface')}")
            sz = el.get("sz")
            if sz is not None and int(sz) < int(MIN_PT * 100):
                probs.add(f"グラフのサイズ{int(sz) / 100}pt(<7pt)")
    for p in sorted(probs):
        v.append((sno, sh.name, "font", p))


def check_text_fit_frame(sno, sh, tf, w_in, h_in, v):
    paras, missing = tf_paras(tf)
    if not any(t for t, _ in paras):
        return
    ml, mr, mt, mb = tf_margins(tf)
    need, lines = measure(paras, w_in - ml - mr)
    avail_h = h_in - mt - mb
    if need > avail_h + 1e-6:
        v.append((sno, sh.name, "text-fit",
                  f"必要 {need:.2f}in > 可用 {avail_h:.2f}in (推定 {lines} 行)"))


def check_text_fit(sno, sh, v):
    x, y, w, h = IN(sh.left), IN(sh.top), IN(sh.width), IN(sh.height)
    if sh.has_table:
        tbl = sh.table
        col_ws = [IN(c.width) for c in tbl.columns]
        for r, row in enumerate(tbl.rows):
            rh = IN(row.height)
            for c, cell in enumerate(row.cells):
                tf = cell.text_frame
                paras, _ = tf_paras(tf)
                if not any(t for t, _ in paras):
                    continue
                ml, mr = IN(cell.margin_left), IN(cell.margin_right)
                mt, mb = IN(cell.margin_top), IN(cell.margin_bottom)
                need, lines = measure(paras, col_ws[c] - ml - mr)
                avail = rh - mt - mb
                if need > avail + 1e-6:
                    v.append((sno, sh.name, "text-fit",
                              f"セル({r},{c}) 必要 {need:.2f}in > 行高の可用 {avail:.2f}in"))
        return
    if sh.has_text_frame:
        check_text_fit_frame(sno, sh, sh.text_frame, w, h, v)


def text_extent(sh):
    """テキストの占有域 (x0, x1, y0, y1)。anchor を考慮した実際の文字域で重なりを見る。"""
    x, y, w, h = IN(sh.left), IN(sh.top), IN(sh.width), IN(sh.height)
    if sh.has_table:
        return x, x + w, y, y + h
    tf = sh.text_frame
    ml, mr, mt, mb = tf_margins(tf)
    paras, _ = tf_paras(tf)
    need, _ = measure(paras, w - ml - mr)
    anchor = tf.vertical_anchor
    avail_h = h - mt - mb
    if anchor == MSO_ANCHOR.BOTTOM:
        top = y + h - mb - need
    elif anchor == MSO_ANCHOR.MIDDLE:
        top = y + mt + (avail_h - need) / 2.0
    else:
        top = y + mt
    return x, x + w, top, top + need


def check_overlap(sno, shapes, v):
    items = [s for s in shapes
             if not s.name.startswith("bg_")
             and (s.has_table or (s.has_text_frame and s.text_frame.text.strip()))]
    ext = {id(s): text_extent(s) for s in items}
    for i, a in enumerate(items):
        for b in items[i + 1:]:
            ax0, ax1, ay0, ay1 = ext[id(a)]
            bx0, bx1, by0, by1 = ext[id(b)]
            ox = min(ax1, bx1) - max(ax0, bx0)
            oy = min(ay1, by1) - max(ay0, by0)
            if ox >= TEXT_TOL and oy >= TEXT_TOL:
                v.append((sno, a.name, "overlap",
                          f"「{b.name}」と重なる (横{ox:.2f}in 縦{oy:.2f}in)"))


def check_title(sno, shapes, v):
    cand = None
    for sh in shapes:
        if (sh.has_text_frame and sh.text_frame.text.strip()
                and abs(IN(sh.left) - 0.5) < 0.02 and abs(IN(sh.top) - 0.25) < 0.02):
            cand = sh
            break
    if cand is None:
        v.append((sno, "-", "title", "タイトル位置(0.5, 0.25)にテキストがない"))
        return
    t = cand.text_frame.text.strip()
    if t.endswith(BANNED_TITLE_END):
        v.append((sno, cand.name, "title", f"体言止めのタイトル: 「{t[-6:]}」で終わる"))


def check_source(sno, shapes, v):
    ok = any(sh.has_text_frame and sh.text_frame.text.strip().startswith("出所:") for sh in shapes)
    if not ok:
        v.append((sno, "-", "source", "「出所:」で始まるテキストがない"))


# ---- 本体 -------------------------------------------------------------------
def audit(path):
    prs = Presentation(path)
    v = []
    for sno, slide in enumerate(prs.slides, 1):
        shapes = list(slide.shapes)
        stype = slide_type(slide)
        for sh in shapes:
            check_bounds(sno, sh, v)
            check_palette(sno, sh, v)
            check_font(sno, sh, v)
            check_text_fit(sno, sh, v)
        check_overlap(sno, shapes, v)
        if stype not in SPECIAL:
            check_title(sno, shapes, v)
            check_source(sno, shapes, v)
    return v


def main():
    if len(sys.argv) != 2:
        print("usage: python3 audit.py deck.pptx", file=sys.stderr)
        return 2
    v = audit(sys.argv[1])
    for sno, name, item, detail in v:
        print(f"{sno} / {name} / {item} / {detail}")
    if not v:
        print("AUDIT PASS")
        return 0
    print(f"AUDIT FAIL: 違反 {len(v)} 件")
    return 1


if __name__ == "__main__":
    sys.exit(main())
